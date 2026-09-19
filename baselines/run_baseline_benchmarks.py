#!/usr/bin/env python
r"""
Phase 11R — FairFullBudget-style baseline rebuild and prediction export.

This is the corrected replacement for the weak Phase 11 V3 baseline recipe.

Core design basis
-----------------
The script follows the stronger historical ResNet18 baseline pathway used in the
Paper 1 project:
- full-budget training rather than the short/over-regularized V3 recipe,
- AdamW with lr=3e-4 and weight_decay=1e-4 by default,
- checkpoint selection using validation balanced accuracy at threshold 0.5,
- post-hoc temperature scaling on validation logits,
- final validation threshold selection under specificity >= 0.91,
- held-out test evaluation and full row-level prediction export.

Models supported
----------------
- resnet18
- efficientnet_b0
- convnext_tiny
- deit_tiny

Data access
-----------
The PCam files are read directly from the six HDF5 files located in:
  <path-to-PCam-data>

This avoids the nested-file/folder and torchvision.PCAM path issues encountered
earlier. HDF5 file locking is disabled for these static read-only benchmark files.

Recommended workflow
--------------------
1. Run ResNet18-only first to confirm this stronger protocol returns to the
   expected historical comparator regime.
2. If the ResNet18 result is acceptable, run all four models sequentially.

The script can also run all four models in one unattended sequence.
"""

from __future__ import annotations

import argparse
import csv
import gc
import json
import math
import os
import random
import sys
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

# Static read-only PCam HDF5 inputs; set before h5py import.
os.environ.setdefault("HDF5_USE_FILE_LOCKING", "FALSE")

import h5py
import numpy as np
import torch
import torch.nn as nn
import torch.nn.functional as F
from PIL import Image
from torch.optim import AdamW
from torch.utils.data import DataLoader, Dataset
from torchvision import models

try:
    from sklearn.metrics import (
        accuracy_score,
        average_precision_score,
        balanced_accuracy_score,
        brier_score_loss,
        cohen_kappa_score,
        confusion_matrix,
        f1_score,
        log_loss,
        matthews_corrcoef,
        precision_score,
        recall_score,
        roc_auc_score,
    )
except Exception as exc:
    raise RuntimeError("scikit-learn is required for Phase 11R metrics.") from exc


SUPPORTED_MODELS = ("resnet18", "efficientnet_b0", "convnext_tiny", "deit_tiny")


# ------------------------------
# General utilities
# ------------------------------

def now() -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S")


def status(msg: str) -> None:
    print(f"[{now()}] {msg}", flush=True)


def jsonable(obj: Any) -> Any:
    if isinstance(obj, dict):
        return {str(k): jsonable(v) for k, v in obj.items()}
    if isinstance(obj, (list, tuple)):
        return [jsonable(v) for v in obj]
    if isinstance(obj, np.ndarray):
        return obj.tolist()
    if isinstance(obj, (np.integer,)):
        return int(obj)
    if isinstance(obj, (np.floating,)):
        return float(obj)
    if isinstance(obj, Path):
        return str(obj)
    return obj


def write_json(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as f:
        json.dump(jsonable(payload), f, indent=2, ensure_ascii=False)


def write_csv(path: Path, rows: Sequence[Dict[str, Any]], fields: Optional[Sequence[str]] = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if fields is None:
        fields = list(rows[0].keys()) if rows else []
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(fields))
        writer.writeheader()
        for row in rows:
            writer.writerow({k: jsonable(row.get(k)) for k in fields})


def try_write_xlsx(path: Path, rows: Sequence[Dict[str, Any]], sheet_name: str = "Summary") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except Exception:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    if not rows:
        ws["A1"] = "No completed runs found."
        wb.save(path)
        return

    fields = list(rows[0].keys())
    ws.append(fields)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([jsonable(row.get(field)) for field in fields])

    for col_idx, field in enumerate(fields, start=1):
        width = min(max(len(str(field)) + 2, 12), 34)
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            width = min(max(width, len(str(value)) + 2), 34)
            if isinstance(value, float):
                ws.cell(row=row_idx, column=col_idx).number_format = "0.0000"
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = "A2"
    wb.save(path)


def seed_everything(seed: int) -> None:
    random.seed(seed)
    np.random.seed(seed)
    torch.manual_seed(seed)
    try:
        torch.cuda.manual_seed_all(seed)
    except Exception:
        pass
    torch.backends.cudnn.deterministic = True
    torch.backends.cudnn.benchmark = False


def resolve_device(device_arg: str) -> torch.device:
    if device_arg == "cpu":
        return torch.device("cpu")
    if device_arg == "cuda":
        if not torch.cuda.is_available():
            raise RuntimeError("--device cuda requested, but CUDA is not available.")
        return torch.device("cuda")
    return torch.device("cuda" if torch.cuda.is_available() else "cpu")


def parameter_count(model: nn.Module) -> int:
    return int(sum(p.numel() for p in model.parameters() if p.requires_grad))


def safe_delete_model(model: Optional[nn.Module], device: torch.device) -> None:
    try:
        del model
    except Exception:
        pass
    gc.collect()
    if device.type == "cuda":
        torch.cuda.empty_cache()


# ------------------------------
# PCam direct HDF5 dataset
# ------------------------------

class PCamDirectH5Dataset(Dataset):
    FILES = {
        "train": (
            "camelyonpatch_level_2_split_train_x.h5",
            "camelyonpatch_level_2_split_train_y.h5",
        ),
        "valid": (
            "camelyonpatch_level_2_split_valid_x.h5",
            "camelyonpatch_level_2_split_valid_y.h5",
        ),
        "test": (
            "camelyonpatch_level_2_split_test_x.h5",
            "camelyonpatch_level_2_split_test_y.h5",
        ),
    }

    def __init__(
        self,
        pcam_dir: str,
        split: str,
        augment: bool = False,
        seed: int = 42,
        limit: Optional[int] = None,
    ):
        if split not in self.FILES:
            raise ValueError(f"Unsupported PCam split: {split}")
        self.pcam_dir = Path(pcam_dir)
        self.split = split
        self.augment = bool(augment and split == "train")
        self.seed = int(seed)
        self.x_path = self.pcam_dir / self.FILES[split][0]
        self.y_path = self.pcam_dir / self.FILES[split][1]
        self._fx = None
        self._fy = None
        self._x = None
        self._y = None

        for path in (self.x_path, self.y_path):
            if not path.exists() or not path.is_file():
                raise FileNotFoundError(f"Expected PCam HDF5 file not found as a regular file: {path}")

        with self._open_h5(self.x_path) as fx:
            n_full = int(fx["x"].shape[0])

        if limit is not None:
            if limit <= 0:
                raise ValueError("Subset limit must be positive.")
            n_limit = min(int(limit), n_full)
            rng = np.random.default_rng(self.seed + {"train": 101, "valid": 202, "test": 303}[split])
            self.indices = np.sort(rng.choice(n_full, size=n_limit, replace=False)).astype(np.int64)
        else:
            self.indices = np.arange(n_full, dtype=np.int64)

    @staticmethod
    def _open_h5(path: Path):
        try:
            return h5py.File(path, "r", locking=False)
        except TypeError:
            return h5py.File(path, "r")

    def _ensure_open(self) -> None:
        if self._fx is None:
            self._fx = self._open_h5(self.x_path)
            self._x = self._fx["x"]
        if self._fy is None:
            self._fy = self._open_h5(self.y_path)
            self._y = self._fy["y"]

    def __len__(self) -> int:
        return int(len(self.indices))

    @staticmethod
    def _augment_chw(x: np.ndarray) -> np.ndarray:
        # Conservative histology-safe augmentation based on the prior Phase 8 export logic.
        if np.random.rand() < 0.5:
            x = x[:, :, ::-1]
        if np.random.rand() < 0.5:
            x = x[:, ::-1, :]
        k = int(np.random.randint(0, 4))
        if k:
            x = np.rot90(x, k=k, axes=(1, 2))
        if np.random.rand() < 0.70:
            gain = np.random.uniform(0.92, 1.08, size=(3, 1, 1)).astype(np.float32)
            bias = np.random.uniform(-0.025, 0.025, size=(3, 1, 1)).astype(np.float32)
            x = x * gain + bias
        return np.ascontiguousarray(np.clip(x, 0.0, 1.0).astype(np.float32))

    def __getitem__(self, local_index: int):
        self._ensure_open()
        raw_index = int(self.indices[int(local_index)])
        x = np.asarray(self._x[raw_index])
        y = float(np.asarray(self._y[raw_index]).reshape(-1)[0])
        if x.ndim == 3 and x.shape[-1] == 3:
            x = np.transpose(x, (2, 0, 1))
        x = x.astype(np.float32) / 255.0
        if self.augment:
            x = self._augment_chw(x)
        else:
            x = np.ascontiguousarray(x)
        return torch.from_numpy(x), torch.tensor(y, dtype=torch.float32), raw_index

    def close(self) -> None:
        for attr in ("_fx", "_fy"):
            handle = getattr(self, attr, None)
            if handle is not None:
                try:
                    handle.close()
                except Exception:
                    pass
                setattr(self, attr, None)
        self._x = None
        self._y = None

    def __del__(self):
        try:
            self.close()
        except Exception:
            pass

    def __getstate__(self):
        self.close()
        return self.__dict__.copy()


# ------------------------------
# Models
# ------------------------------

def build_model(model_name: str) -> nn.Module:
    if model_name == "resnet18":
        model = models.resnet18(weights=None)
        in_features = model.fc.in_features
        model.fc = nn.Linear(in_features, 1)
        return model

    if model_name == "efficientnet_b0":
        model = models.efficientnet_b0(weights=None)
        in_features = model.classifier[-1].in_features
        model.classifier[-1] = nn.Linear(in_features, 1)
        return model

    if model_name == "convnext_tiny":
        model = models.convnext_tiny(weights=None)
        in_features = model.classifier[-1].in_features
        model.classifier[-1] = nn.Linear(in_features, 1)
        return model

    if model_name == "deit_tiny":
        try:
            import timm
        except Exception as exc:
            raise RuntimeError("DeiT-Tiny requires timm. Install with: py -m pip install timm") from exc
        return timm.create_model(
            "deit_tiny_patch16_224",
            pretrained=False,
            img_size=96,
            in_chans=3,
            num_classes=1,
        )

    raise ValueError(f"Unsupported model name: {model_name}")


def logits_from_model(model: nn.Module, x: torch.Tensor) -> torch.Tensor:
    out = model(x)
    if isinstance(out, dict):
        out = out.get("logits", out)
    return out.view(-1)


# ------------------------------
# Metrics / thresholding / calibration
# ------------------------------

def sigmoid_np(logits: np.ndarray) -> np.ndarray:
    logits = np.asarray(logits, dtype=np.float64)
    return 1.0 / (1.0 + np.exp(-np.clip(logits, -60.0, 60.0)))


def confusion_counts(y_true: np.ndarray, pred: np.ndarray) -> Tuple[int, int, int, int]:
    tn, fp, fn, tp = confusion_matrix(y_true, pred, labels=[0, 1]).ravel()
    return int(tp), int(tn), int(fp), int(fn)


def expected_calibration_error(y_true: np.ndarray, prob: np.ndarray, n_bins: int = 15) -> float:
    y_true = np.asarray(y_true, dtype=float)
    prob = np.asarray(prob, dtype=float)
    bins = np.linspace(0.0, 1.0, n_bins + 1)
    ece = 0.0
    for i in range(n_bins):
        lo, hi = bins[i], bins[i + 1]
        if i < n_bins - 1:
            mask = (prob >= lo) & (prob < hi)
        else:
            mask = (prob >= lo) & (prob <= hi)
        if np.any(mask):
            conf = float(np.mean(prob[mask]))
            acc = float(np.mean(y_true[mask]))
            ece += float(np.mean(mask)) * abs(acc - conf)
    return float(ece)


def metrics_from_probs(y_true: np.ndarray, prob: np.ndarray, threshold: float, ece_bins: int = 15) -> Dict[str, Any]:
    y_true = np.asarray(y_true).astype(int)
    prob = np.asarray(prob, dtype=float)
    pred = (prob >= threshold).astype(int)
    tp, tn, fp, fn = confusion_counts(y_true, pred)

    sensitivity = tp / (tp + fn) if (tp + fn) else float("nan")
    specificity = tn / (tn + fp) if (tn + fp) else float("nan")

    return {
        "threshold": float(threshold),
        "roc_auc": float(roc_auc_score(y_true, prob)),
        "pr_auc": float(average_precision_score(y_true, prob)),
        "accuracy": float(accuracy_score(y_true, pred)),
        "balanced_accuracy": float(balanced_accuracy_score(y_true, pred)),
        "positive_f1": float(f1_score(y_true, pred, zero_division=0)),
        "weighted_precision": float(precision_score(y_true, pred, average="weighted", zero_division=0)),
        "weighted_recall": float(recall_score(y_true, pred, average="weighted", zero_division=0)),
        "weighted_f1": float(f1_score(y_true, pred, average="weighted", zero_division=0)),
        "macro_f1": float(f1_score(y_true, pred, average="macro", zero_division=0)),
        "sensitivity": float(sensitivity),
        "specificity": float(specificity),
        "mcc": float(matthews_corrcoef(y_true, pred)),
        "cohen_kappa": float(cohen_kappa_score(y_true, pred)),
        "brier": float(brier_score_loss(y_true, prob)),
        "ece": expected_calibration_error(y_true, prob, n_bins=ece_bins),
        "nll": float(log_loss(y_true, prob, labels=[0, 1])),
        "tp": tp,
        "tn": tn,
        "fp": fp,
        "fn": fn,
    }


def threshold_sweep(
    y_true: np.ndarray,
    prob: np.ndarray,
    min_specificity: float,
    threshold_step: float,
) -> Tuple[float, str, Dict[str, Any], List[Dict[str, Any]]]:
    thresholds = np.round(np.arange(0.005, 0.995 + threshold_step / 2, threshold_step), 6)
    rows: List[Dict[str, Any]] = []
    for t in thresholds:
        m = metrics_from_probs(y_true, prob, float(t))
        rows.append(m)

    feasible = [row for row in rows if row["specificity"] >= min_specificity]
    if feasible:
        pool = feasible
        selection_type = "specificity_constrained_balanced_accuracy"
    else:
        pool = rows
        selection_type = "fallback_unconstrained_balanced_accuracy"

    best = max(
        pool,
        key=lambda row: (
            float(row["balanced_accuracy"]),
            float(row["positive_f1"]),
            float(row["sensitivity"]),
            -float(row["threshold"]),
        ),
    )
    return float(best["threshold"]), selection_type, best, rows


def fit_temperature(validation_logits: np.ndarray, validation_y: np.ndarray, device: torch.device) -> float:
    logits = torch.as_tensor(validation_logits, dtype=torch.float32, device=device)
    labels = torch.as_tensor(validation_y, dtype=torch.float32, device=device)
    log_temp = torch.zeros(1, dtype=torch.float32, device=device, requires_grad=True)
    optimizer = torch.optim.LBFGS([log_temp], lr=0.1, max_iter=50)

    def closure():
        optimizer.zero_grad()
        temp = torch.exp(log_temp).clamp(min=1e-3, max=1e3)
        loss = F.binary_cross_entropy_with_logits(logits / temp, labels)
        loss.backward()
        return loss

    try:
        optimizer.step(closure)
        temp = float(torch.exp(log_temp).detach().cpu().item())
        if not math.isfinite(temp) or temp <= 0:
            return 1.0
        return temp
    except Exception:
        return 1.0


# ------------------------------
# Train/evaluate
# ------------------------------

@dataclass
class CollectedLogits:
    indices: np.ndarray
    y_true: np.ndarray
    logits: np.ndarray
    loss: float
    elapsed_seconds: float


@torch.no_grad()
def collect_logits(
    model: nn.Module,
    loader: DataLoader,
    device: torch.device,
) -> CollectedLogits:
    model.eval()
    loss_fn = nn.BCEWithLogitsLoss()
    all_indices: List[np.ndarray] = []
    all_y: List[np.ndarray] = []
    all_logits: List[np.ndarray] = []
    total_loss = 0.0
    total_n = 0
    t0 = time.perf_counter()
    for x, y, idx in loader:
        x = x.to(device, non_blocking=(device.type == "cuda"))
        y = y.to(device, non_blocking=(device.type == "cuda")).float().view(-1)
        logits = logits_from_model(model, x)
        loss = loss_fn(logits, y)
        n = int(y.numel())
        total_loss += float(loss.detach().cpu().item()) * n
        total_n += n
        all_indices.append(np.asarray(idx.detach().cpu().numpy(), dtype=np.int64))
        all_y.append(np.asarray(y.detach().cpu().numpy(), dtype=np.int64))
        all_logits.append(np.asarray(logits.detach().cpu().numpy(), dtype=np.float64))
    elapsed = time.perf_counter() - t0
    return CollectedLogits(
        indices=np.concatenate(all_indices),
        y_true=np.concatenate(all_y).astype(int),
        logits=np.concatenate(all_logits),
        loss=total_loss / max(1, total_n),
        elapsed_seconds=float(elapsed),
    )


def train_one_epoch(
    model: nn.Module,
    loader: DataLoader,
    optimizer: AdamW,
    device: torch.device,
    grad_clip: float,
) -> float:
    model.train()
    loss_fn = nn.BCEWithLogitsLoss()
    total_loss = 0.0
    total_n = 0
    for x, y, _idx in loader:
        x = x.to(device, non_blocking=(device.type == "cuda"))
        y = y.to(device, non_blocking=(device.type == "cuda")).float().view(-1)
        optimizer.zero_grad(set_to_none=True)
        logits = logits_from_model(model, x)
        loss = loss_fn(logits, y)
        loss.backward()
        if grad_clip > 0:
            nn.utils.clip_grad_norm_(model.parameters(), grad_clip)
        optimizer.step()
        n = int(y.numel())
        total_loss += float(loss.detach().cpu().item()) * n
        total_n += n
    return total_loss / max(1, total_n)


def save_checkpoint(
    path: Path,
    model: nn.Module,
    epoch: int,
    valid_default_metrics: Dict[str, Any],
    config: Dict[str, Any],
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    torch.save(
        {
            "epoch": int(epoch),
            "model_state_dict": model.state_dict(),
            "valid_default_metrics": jsonable(valid_default_metrics),
            "config": jsonable(config),
        },
        path,
    )


def export_prediction_csv(
    path: Path,
    split: str,
    model_name: str,
    indices: np.ndarray,
    y_true: np.ndarray,
    logits: np.ndarray,
    prob_uncalibrated: np.ndarray,
    prob_calibrated: np.ndarray,
    selected_threshold: float,
    temperature: float,
    checkpoint_path: Path,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rows: List[Dict[str, Any]] = []
    pred_selected = (prob_calibrated >= selected_threshold).astype(int)
    pred_cal_0p5 = (prob_calibrated >= 0.5).astype(int)
    pred_uncal_0p5 = (prob_uncalibrated >= 0.5).astype(int)
    for i in range(len(y_true)):
        rows.append(
            {
                "split": split,
                "index": int(indices[i]),
                "y_true": int(y_true[i]),
                "logit": float(logits[i]),
                "probability_uncalibrated": float(prob_uncalibrated[i]),
                "probability_calibrated": float(prob_calibrated[i]),
                "pred_selected": int(pred_selected[i]),
                "pred_calibrated_0p5": int(pred_cal_0p5[i]),
                "pred_uncalibrated_0p5": int(pred_uncal_0p5[i]),
                "selected_threshold": float(selected_threshold),
                "temperature": float(temperature),
                "model": model_name,
                "checkpoint": str(checkpoint_path),
            }
        )
    write_csv(path, rows)


def category_counts(y_true: np.ndarray, prob: np.ndarray, threshold: float) -> Dict[str, int]:
    pred = (prob >= threshold).astype(int)
    tp, tn, fp, fn = confusion_counts(y_true, pred)
    return {"tp": tp, "tn": tn, "fp": fp, "fn": fn}


def build_loaders(args: argparse.Namespace, model_seed: int) -> Tuple[DataLoader, DataLoader, DataLoader, Dict[str, int]]:
    train_ds = PCamDirectH5Dataset(
        args.pcam_dir, "train", augment=(not args.no_augment), seed=model_seed, limit=args.train_limit
    )
    valid_ds = PCamDirectH5Dataset(
        args.pcam_dir, "valid", augment=False, seed=model_seed, limit=args.valid_limit
    )
    test_ds = PCamDirectH5Dataset(
        args.pcam_dir, "test", augment=False, seed=model_seed, limit=args.test_limit
    )
    pin_memory = torch.cuda.is_available() and args.device != "cpu"
    train_loader = DataLoader(
        train_ds,
        batch_size=args.batch_size,
        shuffle=True,
        num_workers=args.num_workers,
        pin_memory=pin_memory,
        drop_last=False,
        persistent_workers=(args.num_workers > 0),
    )
    valid_loader = DataLoader(
        valid_ds,
        batch_size=args.batch_size,
        shuffle=False,
        num_workers=args.num_workers,
        pin_memory=pin_memory,
        drop_last=False,
        persistent_workers=(args.num_workers > 0),
    )
    test_loader = DataLoader(
        test_ds,
        batch_size=args.batch_size,
        shuffle=False,
        num_workers=args.num_workers,
        pin_memory=pin_memory,
        drop_last=False,
        persistent_workers=(args.num_workers > 0),
    )
    counts = {"train": len(train_ds), "valid": len(valid_ds), "test": len(test_ds)}
    return train_loader, valid_loader, test_loader, counts


def run_one_model(args: argparse.Namespace, model_name: str, model_position: int, device: torch.device) -> Dict[str, Any]:
    model_outdir = Path(args.base_outdir) / model_name
    model_outdir.mkdir(parents=True, exist_ok=True)
    pred_dir = model_outdir / "predictions"
    ckpt_dir = model_outdir / "checkpoints"
    table_dir = model_outdir / "tables"
    for p in (pred_dir, ckpt_dir, table_dir):
        p.mkdir(parents=True, exist_ok=True)

    summary_path = model_outdir / "summary.json"
    pred_test_path = pred_dir / "test_predictions.csv"
    if args.skip_completed and summary_path.exists() and pred_test_path.exists():
        status(f"Skipping completed model={model_name}: summary and test predictions already exist.")
        with summary_path.open("r", encoding="utf-8") as f:
            return json.load(f)

    model_seed = int(args.seed + 1000 * model_position)
    seed_everything(model_seed)
    train_loader, valid_loader, test_loader, counts = build_loaders(args, model_seed)

    model = build_model(model_name).to(device)
    n_params = parameter_count(model)
    optimizer = AdamW(model.parameters(), lr=args.lr, weight_decay=args.weight_decay)

    config = {
        "phase": "Phase 11R FairFullBudget-style all-baseline rebuild",
        "model": model_name,
        "pcam_dir": args.pcam_dir,
        "outdir": str(model_outdir),
        "device": str(device),
        "epochs": args.epochs,
        "batch_size": args.batch_size,
        "num_workers": args.num_workers,
        "lr": args.lr,
        "weight_decay": args.weight_decay,
        "seed": model_seed,
        "base_seed": args.seed,
        "min_specificity": args.min_specificity,
        "threshold_step": args.threshold_step,
        "ece_bins": args.ece_bins,
        "grad_clip": args.grad_clip,
        "augment_train": not args.no_augment,
        "checkpoint_selection": "best validation balanced accuracy at default threshold 0.5",
        "posthoc_calibration": "validation temperature scaling",
        "final_threshold_selection": "validation calibrated probabilities; maximize balanced accuracy subject to specificity >= min_specificity",
        "training_from_scratch": True,
        "parameter_count": n_params,
        "dataset_counts": counts,
    }
    write_json(model_outdir / "run_config.json", config)

    status(
        f"Starting model={model_name}, device={device}, parameters={n_params:,}, "
        f"epochs={args.epochs}, dataset_counts={counts}"
    )

    history: List[Dict[str, Any]] = []
    best_bacc = -1.0
    best_epoch = None
    best_ckpt = ckpt_dir / "best_valid_balanced_accuracy.pt"
    run_start = time.perf_counter()

    for epoch in range(1, args.epochs + 1):
        epoch_start = time.perf_counter()
        train_loss = train_one_epoch(model, train_loader, optimizer, device, grad_clip=args.grad_clip)
        valid_bundle = collect_logits(model, valid_loader, device)
        valid_prob_uncal = sigmoid_np(valid_bundle.logits)
        valid_default = metrics_from_probs(valid_bundle.y_true, valid_prob_uncal, threshold=0.5, ece_bins=args.ece_bins)
        valid_default["loss"] = float(valid_bundle.loss)
        row = {
            "epoch": epoch,
            "train_loss": float(train_loss),
            "epoch_seconds": float(time.perf_counter() - epoch_start),
        }
        for k, v in valid_default.items():
            row[f"valid_default_{k}"] = v
        history.append(row)
        write_json(model_outdir / "history.json", history)
        write_csv(model_outdir / "history.csv", history)

        current_bacc = float(valid_default["balanced_accuracy"])
        marker = ""
        if current_bacc > best_bacc:
            best_bacc = current_bacc
            best_epoch = int(epoch)
            save_checkpoint(best_ckpt, model, epoch, valid_default, config)
            marker = " [NEW BEST]"
        status(
            f"model={model_name} epoch={epoch}/{args.epochs} "
            f"train_loss={train_loss:.6f} valid_auc={valid_default['roc_auc']:.6f} "
            f"valid_pr={valid_default['pr_auc']:.6f} valid_bacc={valid_default['balanced_accuracy']:.6f}"
            f"{marker}"
        )

    if best_epoch is None:
        raise RuntimeError(f"No best checkpoint selected for model={model_name}")

    # Reload selected checkpoint
    checkpoint = torch.load(best_ckpt, map_location=device)
    model.load_state_dict(checkpoint["model_state_dict"])
    model.eval()

    valid_bundle = collect_logits(model, valid_loader, device)
    test_bundle = collect_logits(model, test_loader, device)

    valid_prob_uncal = sigmoid_np(valid_bundle.logits)
    test_prob_uncal = sigmoid_np(test_bundle.logits)

    temperature = fit_temperature(valid_bundle.logits, valid_bundle.y_true, device)
    valid_prob_cal = sigmoid_np(valid_bundle.logits / temperature)
    test_prob_cal = sigmoid_np(test_bundle.logits / temperature)

    selected_threshold, threshold_type, valid_selected, sweep_rows = threshold_sweep(
        valid_bundle.y_true,
        valid_prob_cal,
        min_specificity=args.min_specificity,
        threshold_step=args.threshold_step,
    )
    write_csv(table_dir / "validation_threshold_sweep_calibrated.csv", sweep_rows)

    valid_uncal_default = metrics_from_probs(valid_bundle.y_true, valid_prob_uncal, 0.5, args.ece_bins)
    valid_uncal_default["loss"] = float(valid_bundle.loss)
    valid_cal_default = metrics_from_probs(valid_bundle.y_true, valid_prob_cal, 0.5, args.ece_bins)
    valid_cal_default["loss"] = float(valid_bundle.loss)
    valid_selected["loss"] = float(valid_bundle.loss)

    test_uncal_default = metrics_from_probs(test_bundle.y_true, test_prob_uncal, 0.5, args.ece_bins)
    test_uncal_default["loss"] = float(test_bundle.loss)
    test_cal_default = metrics_from_probs(test_bundle.y_true, test_prob_cal, 0.5, args.ece_bins)
    test_cal_default["loss"] = float(test_bundle.loss)
    test_selected = metrics_from_probs(test_bundle.y_true, test_prob_cal, selected_threshold, args.ece_bins)
    test_selected["loss"] = float(test_bundle.loss)

    export_prediction_csv(
        pred_dir / "validation_predictions.csv",
        split="valid",
        model_name=model_name,
        indices=valid_bundle.indices,
        y_true=valid_bundle.y_true,
        logits=valid_bundle.logits,
        prob_uncalibrated=valid_prob_uncal,
        prob_calibrated=valid_prob_cal,
        selected_threshold=selected_threshold,
        temperature=temperature,
        checkpoint_path=best_ckpt,
    )
    export_prediction_csv(
        pred_test_path,
        split="test",
        model_name=model_name,
        indices=test_bundle.indices,
        y_true=test_bundle.y_true,
        logits=test_bundle.logits,
        prob_uncalibrated=test_prob_uncal,
        prob_calibrated=test_prob_cal,
        selected_threshold=selected_threshold,
        temperature=temperature,
        checkpoint_path=best_ckpt,
    )

    scalar_rows: List[Dict[str, Any]] = []
    for split_name, mode_name, metrics_dict in (
        ("valid", "uncalibrated_default_0p5", valid_uncal_default),
        ("valid", "calibrated_default_0p5", valid_cal_default),
        ("valid", "calibrated_selected_threshold", valid_selected),
        ("test", "uncalibrated_default_0p5", test_uncal_default),
        ("test", "calibrated_default_0p5", test_cal_default),
        ("test", "calibrated_selected_threshold", test_selected),
    ):
        row = {"split": split_name, "mode": mode_name}
        row.update(metrics_dict)
        scalar_rows.append(row)
    write_csv(table_dir / "scalar_metrics.csv", scalar_rows)

    elapsed = time.perf_counter() - run_start
    inference_ms_per_patch = 1000.0 * test_bundle.elapsed_seconds / max(1, len(test_bundle.y_true))

    summary = {
        "phase": "Phase 11R FairFullBudget-style all-baseline rebuild",
        "model": model_name,
        "parameter_count": n_params,
        "training_from_scratch": True,
        "dataset_counts": counts,
        "best_epoch": int(best_epoch),
        "checkpoint_selection": "best validation balanced accuracy at threshold 0.5",
        "checkpoint": str(best_ckpt),
        "temperature": float(temperature),
        "selected_threshold": float(selected_threshold),
        "selected_threshold_type": threshold_type,
        "min_specificity": float(args.min_specificity),
        "validation": {
            "uncalibrated_default_0p5": valid_uncal_default,
            "calibrated_default_0p5": valid_cal_default,
            "calibrated_selected_threshold": valid_selected,
        },
        "test": {
            "uncalibrated_default_0p5": test_uncal_default,
            "calibrated_default_0p5": test_cal_default,
            "calibrated_selected_threshold": test_selected,
        },
        "prediction_exports": {
            "validation": str(pred_dir / "validation_predictions.csv"),
            "test": str(pred_test_path),
        },
        "tables": {
            "threshold_sweep_calibrated_validation": str(table_dir / "validation_threshold_sweep_calibrated.csv"),
            "scalar_metrics": str(table_dir / "scalar_metrics.csv"),
        },
        "inference_seconds_test_split": float(test_bundle.elapsed_seconds),
        "inference_ms_per_patch_test_split": float(inference_ms_per_patch),
        "elapsed_seconds_total": float(elapsed),
        "run_config": config,
    }
    write_json(summary_path, summary)
    write_json(model_outdir / "final_selected_test_metrics.json", {
        "model": model_name,
        "parameter_count": n_params,
        "best_epoch": int(best_epoch),
        "selected_threshold": float(selected_threshold),
        "temperature": float(temperature),
        **test_selected,
        "inference_ms_per_patch_test_split": float(inference_ms_per_patch),
    })

    status(
        f"Completed model={model_name}. "
        f"selected_test_auc={test_selected['roc_auc']:.6f}, "
        f"selected_test_pr={test_selected['pr_auc']:.6f}, "
        f"selected_test_bacc={test_selected['balanced_accuracy']:.6f}, "
        f"selected_test_spec={test_selected['specificity']:.6f}."
    )
    safe_delete_model(model, device)
    return summary


def flatten_summary(summary: Dict[str, Any]) -> Dict[str, Any]:
    test = summary["test"]["calibrated_selected_threshold"]
    return {
        "model": summary["model"],
        "parameter_count": summary["parameter_count"],
        "best_epoch": summary["best_epoch"],
        "temperature": summary["temperature"],
        "selected_threshold": summary["selected_threshold"],
        "roc_auc": test["roc_auc"],
        "pr_auc": test["pr_auc"],
        "accuracy": test["accuracy"],
        "balanced_accuracy": test["balanced_accuracy"],
        "positive_f1": test["positive_f1"],
        "weighted_precision": test["weighted_precision"],
        "weighted_recall": test["weighted_recall"],
        "weighted_f1": test["weighted_f1"],
        "macro_f1": test["macro_f1"],
        "sensitivity": test["sensitivity"],
        "specificity": test["specificity"],
        "brier": test["brier"],
        "ece": test["ece"],
        "mcc": test["mcc"],
        "cohen_kappa": test["cohen_kappa"],
        "nll": test["nll"],
        "tp": test["tp"],
        "tn": test["tn"],
        "fp": test["fp"],
        "fn": test["fn"],
        "inference_ms_per_patch_test_split": summary["inference_ms_per_patch_test_split"],
        "summary_json": str(Path(summary["run_config"]["outdir"]) / "summary.json"),
        "test_prediction_csv": summary["prediction_exports"]["test"],
    }


def collect_all_summaries(base_outdir: Path, model_names: Sequence[str]) -> Dict[str, Any]:
    summary_dir = base_outdir / "summary"
    summary_dir.mkdir(parents=True, exist_ok=True)
    rows: List[Dict[str, Any]] = []
    missing: List[str] = []
    for model_name in model_names:
        path = base_outdir / model_name / "summary.json"
        if not path.exists():
            missing.append(str(path))
            continue
        with path.open("r", encoding="utf-8") as f:
            rows.append(flatten_summary(json.load(f)))
    write_csv(summary_dir / "all_baseline_comparison.csv", rows)
    try_write_xlsx(summary_dir / "all_baseline_comparison.xlsx", rows, sheet_name="Baselines")

    md_path = summary_dir / "all_baseline_comparison.md"
    if rows:
        display_fields = [
            "model", "parameter_count", "best_epoch", "selected_threshold",
            "roc_auc", "pr_auc", "balanced_accuracy", "positive_f1",
            "sensitivity", "specificity", "weighted_f1", "mcc"
        ]
        lines = [
            "| " + " | ".join(display_fields) + " |",
            "| " + " | ".join(["---"] * len(display_fields)) + " |",
        ]
        for row in rows:
            vals = []
            for field in display_fields:
                value = row[field]
                if isinstance(value, float):
                    vals.append(f"{value:.4f}")
                else:
                    vals.append(str(value))
            lines.append("| " + " | ".join(vals) + " |")
        md_path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    else:
        md_path.write_text("No completed Phase 11R baseline summaries were found.\n", encoding="utf-8")

    report = {
        "base_outdir": str(base_outdir),
        "completed_models": [row["model"] for row in rows],
        "missing_summary_paths": missing,
        "csv": str(summary_dir / "all_baseline_comparison.csv"),
        "xlsx_if_openpyxl_available": str(summary_dir / "all_baseline_comparison.xlsx"),
        "markdown": str(md_path),
    }
    write_json(summary_dir / "summary_collection_report.json", report)
    return report


# ------------------------------
# CLI
# ------------------------------

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Phase 11R fair/full-budget baseline runner with calibrated prediction exports."
    )
    parser.add_argument(
        "--pcam-dir",
        required=True,
        help="Folder directly containing the six PCam HDF5 files.",
    )
    parser.add_argument(
        "--base-outdir",
        default=str(Path(__file__).resolve().parents[1] / "results" / "baseline_benchmarks"),
        help="Parent output directory for all Phase 11R runs.",
    )
    parser.add_argument(
        "--models",
        nargs="+",
        default=list(SUPPORTED_MODELS),
        choices=SUPPORTED_MODELS,
        help="Models to run sequentially.",
    )
    parser.add_argument("--epochs", type=int, default=10)
    parser.add_argument("--batch-size", type=int, default=128)
    parser.add_argument("--num-workers", type=int, default=0)
    parser.add_argument("--lr", type=float, default=3e-4)
    parser.add_argument("--weight-decay", type=float, default=1e-4)
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--device", choices=("auto", "cpu", "cuda"), default="auto")
    parser.add_argument("--min-specificity", type=float, default=0.91)
    parser.add_argument("--threshold-step", type=float, default=0.005)
    parser.add_argument("--ece-bins", type=int, default=15)
    parser.add_argument("--grad-clip", type=float, default=5.0)
    parser.add_argument("--no-augment", action="store_true")
    parser.add_argument("--skip-completed", action="store_true")
    parser.add_argument("--continue-on-error", action="store_true")
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--train-limit", type=int, default=None)
    parser.add_argument("--valid-limit", type=int, default=None)
    parser.add_argument("--test-limit", type=int, default=None)
    return parser.parse_args()


def check_only(args: argparse.Namespace, device: torch.device) -> int:
    base_outdir = Path(args.base_outdir)
    base_outdir.mkdir(parents=True, exist_ok=True)
    report: Dict[str, Any] = {
        "check_only": True,
        "device": str(device),
        "pcam_dir": args.pcam_dir,
        "models": [],
    }
    for pos, model_name in enumerate(args.models):
        seed = args.seed + pos * 1000
        train_loader, valid_loader, test_loader, counts = build_loaders(args, seed)
        model = build_model(model_name).to(device)
        info = {
            "model": model_name,
            "parameter_count": parameter_count(model),
            "dataset_counts": counts,
        }
        report["models"].append(info)
        status(f"Check passed: {info}")
        safe_delete_model(model, device)
    write_json(base_outdir / "check_only_report.json", report)
    print(json.dumps(report, indent=2), flush=True)
    return 0


def main() -> int:
    args = parse_args()
    base_outdir = Path(args.base_outdir)
    base_outdir.mkdir(parents=True, exist_ok=True)
    device = resolve_device(args.device)

    if args.check_only:
        return check_only(args, device)

    manifest = {
        "phase": "Phase 11R FairFullBudget-style all-baseline rebuild",
        "start_time": now(),
        "args": vars(args),
        "device": str(device),
        "completed_models": [],
        "failed_models": [],
    }
    write_json(base_outdir / "baseline_run_manifest.json", manifest)

    summaries: List[Dict[str, Any]] = []
    for pos, model_name in enumerate(args.models):
        try:
            summary = run_one_model(args, model_name, pos, device)
            summaries.append(summary)
            manifest["completed_models"].append(model_name)
            write_json(base_outdir / "baseline_run_manifest.json", manifest)
        except Exception as exc:
            failure = {
                "model": model_name,
                "error_type": type(exc).__name__,
                "error_message": str(exc),
                "traceback": traceback.format_exc(),
            }
            failure_path = base_outdir / model_name / "FAILED_run_error.json"
            write_json(failure_path, failure)
            manifest["failed_models"].append(failure)
            write_json(base_outdir / "baseline_run_manifest.json", manifest)
            status(f"FAILED model={model_name}: {type(exc).__name__}: {exc}")
            if not args.continue_on_error:
                raise

    collection_report = collect_all_summaries(base_outdir, args.models)
    manifest["end_time"] = now()
    manifest["summary_collection"] = collection_report
    write_json(base_outdir / "baseline_run_manifest.json", manifest)
    status("Phase 11R sequence completed.")
    print(json.dumps(collection_report, indent=2), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

